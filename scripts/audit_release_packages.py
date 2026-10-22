#!/usr/bin/env python3
"""Audit the public and private release packages before any repository work.

Checks (Section 14 of the governing prompt):
  * no credential values anywhere in either package;
  * no secret-bearing file or directory names in either package;
  * required deliverables are present;
  * manuscript consistency: abstract length, highlights, citations, refs,
    placeholders, claim-evidence status;
  * public package contains only shareable material.

Writes reports/FINAL_AUDIT_REPORT.md. Never prints credential values.
"""
from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PUB = ROOT / "reproducibility" / "public_repo"
PRIV = ROOT / "reproducibility" / "github"
APIS = ROOT / "apis"
REPORT = ROOT / "reports" / "FINAL_AUDIT_REPORT.md"

SECRET_PATTERNS = [
    re.compile(rb"sk-[A-Za-z0-9_\-]{20,}"),
    re.compile(rb"ghp_[A-Za-z0-9]{30,}"),
    re.compile(rb"github_pat_[A-Za-z0-9_]{30,}"),
    re.compile(rb"AKIA[0-9A-Z]{16}"),
    re.compile(rb"xox[baprs]-[A-Za-z0-9-]{10,}"),
    re.compile(rb"-----BEGIN [A-Z ]*PRIVATE KEY-----"),
    re.compile(rb"AIza[0-9A-Za-z_\-]{30,}"),
]
SECRET_NAMES = re.compile(r"(^|/)(apis|\.env|\.git|secrets?)(/|$)|\.(key|pem)$", re.I)


CREDENTIAL_FILE = re.compile(r"(chatgpt|deepseek|ncbi|openalex|github_fine_grained|meta api|open access api)", re.I)


def secret_tokens() -> tuple[set[bytes], list[str]]:
    tokens: set[bytes] = set()
    names = []
    if not APIS.exists():
        return tokens, names
    for path in APIS.rglob("*"):
        if not path.is_file():
            continue
        names.append(path.name)
        if not CREDENTIAL_FILE.search(path.name):
            continue
        raw = path.read_bytes()
        for piece in re.split(rb"[\s,;'\"]+", raw):
            piece = piece.strip()
            if len(piece) >= 12:
                tokens.add(piece)
        stripped = raw.strip()
        if len(stripped) >= 12:
            tokens.add(stripped)
    return tokens, names


def scan_tree(root: Path, tokens: set[bytes]) -> dict:
    result = {"files": 0, "secret_value_hits": [], "secret_name_hits": [],
              "pattern_hits": [], "bytes": 0}
    if not root.exists():
        return result
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        result["files"] += 1
        rel = str(path.relative_to(root))
        result["bytes"] += path.stat().st_size
        if SECRET_NAMES.search("/" + rel):
            result["secret_name_hits"].append(rel)
        data = path.read_bytes()
        if any(tok in data for tok in tokens):
            result["secret_value_hits"].append(rel)
        for pat in SECRET_PATTERNS:
            if pat.search(data):
                result["pattern_hits"].append(f"{rel}: {pat.pattern!r}")
    return result


def git_unignored_files() -> list[str]:
    try:
        out = subprocess.run(
            ["git", "ls-files", "--others", "--exclude-standard", "--cached"],
            cwd=ROOT, capture_output=True, text=True, check=False)
        return [line for line in out.stdout.splitlines() if line.strip()]
    except Exception:
        return []


def scan_git_unignored(tokens: set[bytes]) -> dict:
    result = {"files": 0, "secret_value_hits": [], "pattern_hits": []}
    for rel in git_unignored_files():
        path = ROOT / rel
        if not path.is_file() or "apis" in Path(rel).parts:
            continue
        result["files"] += 1
        data = path.read_bytes()
        if any(tok in data for tok in tokens):
            result["secret_value_hits"].append(rel)
        for pat in SECRET_PATTERNS:
            if pat.search(data):
                result["pattern_hits"].append(f"{rel}: {pat.pattern!r}")
    return result


def manuscript_checks() -> dict:
    text = (ROOT / "manuscript" / "Manuscript.md").read_text()
    supp = (ROOT / "manuscript" / "Supplementary_Materials.md").read_text()
    out: dict = {}
    abstract = re.search(r"## Abstract\s+(.*?)\n\n", text, re.S)
    out["abstract_words"] = len(abstract.group(1).split()) if abstract else -1
    hl_block = re.search(r"## Highlights\s+(.*?)\n\n", text, re.S)
    highlights = [l.strip("- ").strip() for l in hl_block.group(1).splitlines() if l.strip().startswith("-")] if hl_block else []
    out["highlights"] = len(highlights)
    out["highlights_max_chars"] = max((len(h) for h in highlights), default=0)
    body, refs = text.split("## References", 1)
    body_clean = re.sub(r"\$\$.*?\$\$", " ", body, flags=re.S)
    body_clean = re.sub(r"\$[^$\n]*\$", " ", body_clean)
    cited: set[int] = set()
    for m in re.finditer(r"\[([0-9,\s\-\u2013\u2014]+)\]", body_clean):
        for part in re.split(r"[,\s]+", m.group(1).strip()):
            if not part:
                continue
            rng = re.match(r"^(\d+)\s*[-\u2013\u2014]{1,2}\s*(\d+)$", part)
            if rng:
                cited.update(range(int(rng.group(1)), int(rng.group(2)) + 1))
            elif part.isdigit():
                cited.add(int(part))
    refnums = {int(m.group(1)) for m in re.finditer(r"^(\d+)\.\s", refs, re.M)}
    out["references"] = len(refnums)
    out["cited"] = len(cited)
    out["cited_missing"] = sorted(cited - refnums)
    out["uncited"] = sorted(refnums - cited)
    out["refs_with_doi"] = len(re.findall(r"doi:10\.", refs))
    out["placeholders"] = re.findall(r"\[[^\]]*(?:to be inserted|pending|TODO|TBD)[^\]]*\]", text, re.I)
    out["raw_display_math_in_docx_source"] = text.count("$$")
    out["supp_has_dimensionless_section"] = "S3.1" in supp
    out["supp_ansatz_exponential"] = "2\\exp" in supp and "s\\to2" in supp and "c_{\\pm}\\to1" in supp
    return out


def claim_checks() -> dict:
    path = ROOT / "manuscript" / "claim_evidence.csv"
    rows = list(csv.DictReader(path.open()))
    quant = [r for r in rows if r.get("claim_id", "").startswith("C")]
    lit = [r for r in rows if r.get("claim_id", "").startswith("L")]
    supported = [r for r in rows if (r.get("status") or r.get("support_status")) == "supported"]
    return {"claims": len(rows), "quant": len(quant), "lit": len(lit),
            "statuses": sorted({(r.get("status") or r.get("support_status") or "") for r in rows}),
            "all_supported": len(supported) == len(rows) and len(rows) > 0,
            "all_resolved": all((r.get("status") or r.get("support_status")) in {"supported", "superseded"} for r in rows) and len(rows) > 0,
            "quant_complete": [f"C{i:02d}" for i in range(1, 45)] == [r["claim_id"] for r in quant],
            "lit_complete": [f"L{i:02d}" for i in range(1, 63)] == [r["claim_id"] for r in lit],
            "all_evidenced": all(r.get("artifact") for r in rows)}


def display_item_checks() -> dict:
    text = (ROOT / "manuscript" / "Manuscript.md").read_text()
    supp = (ROOT / "manuscript" / "Supplementary_Materials.md").read_text()
    body = text.split("## References")[0]
    # Count only in-text citations: exclude figure image blocks and caption lines,
    # which now sit inline immediately after the paragraph that cites them.
    prose = "\n".join(
        l for l in body.splitlines()
        if not l.lstrip().startswith("![")
        and not re.match(r"^\*\*Figure \d+\.\*\*", l.strip())
    )
    order: list[int] = []
    for m in re.finditer(r"Figure\s+(\d+)", prose):
        n = int(m.group(1))
        if n not in order:
            order.append(n)
    body_tables = len(re.findall(r"^\s*\|.*\|\s*$", body, re.M)) + len(re.findall(r"^\*\*Table \d", body, re.M))
    supp_tables = sorted({int(m.group(1)) for m in re.finditer(r"\*\*Supplementary Table S(\d+)\.", supp)})
    cited: set[int] = set()
    for m in re.finditer(r"(?:Supplementary\s+)?Tables?\s+S(\d+(?:\s*(?:,|and|--|\u2013|-)\s*S?\d+)*)", text):
        s = m.group(1).replace("\u2013", "-").replace("--", "-").replace("S", "")
        for part in re.split(r",|;|\band\b", s):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                a, b = part.split("-")[0], part.split("-")[-1]
                cited.update(range(int(a), int(b) + 1))
            else:
                cited.add(int(part))
    return {"figure_order": order, "body_tables": body_tables,
            "supp_tables": len(supp_tables),
            "uncited_supp_tables": [n for n in supp_tables if n not in cited]}


def sequential_checks() -> dict:
    import enforce_sequential_numbering_and_inline_figures as E

    man = E.MAN.read_text()
    sup = E.SUP.read_text()
    body = E.read_doc(E.MAN)[0]
    return {
        "ref_order": E.first_appearance(body, E.REF_RE, "citation"),
        "tab_order": E.first_appearance(body, E.TAB_RE, "table"),
        "eq_main": E.check_equations(man, E.EQ_MAIN_RE),
        "eq_supp": E.check_equations(sup, E.EQ_SUP_RE),
        "figs": E.check_figures_markdown(man),
        "figs_docx": E.check_figures_docx(),
    }


def stage11_checks() -> dict:
    tables = ["table_sgfv_fisher_validation.csv",
              "table_constraint_schemes_4way.csv",
              "table_physics_weight_sensitivity.csv"]
    have_tables = all((ROOT / "results" / "tables" / t).exists() for t in tables)
    upgrades = ROOT / "results" / "stage11_upgrades.json"
    keys = {}
    if upgrades.exists():
        data = json.loads(upgrades.read_text())
        keys = {k: k in data for k in ("fisher_sgfv_validation",
                                       "constraint_schemes_benchmark",
                                       "physics_weight_sweep")}
    stats = json.loads((ROOT / "results" / "statistics.json").read_text())
    stats_keys = all(k in stats for k in ("stage11_fisher_sgfv_validation",
                                          "stage11_constraint_schemes_benchmark",
                                          "stage11_physics_weight_sweep"))
    script = (ROOT / "scripts" / "run_stage11_reviewer_upgrades.py").exists()
    return {"tables": have_tables, "json_keys": keys,
            "statistics_keys": stats_keys, "script": script}


def zendo_doi_check() -> bool:
    import zipfile
    doi = "10.5281/zenodo.23487404"
    md = (ROOT / "manuscript" / "Manuscript.md").read_text()
    readme = (PUB / "README.md").read_text()
    letter = (ROOT / "manuscript" / "Cover_Letter_CPC.docx")
    main = ROOT / "manuscript" / "Manuscript_PINO_clean.docx"
    if doi not in md or doi not in readme or not letter.exists() or not main.exists():
        return False
    for path in (main, letter):
        with zipfile.ZipFile(path) as zf:
            if doi not in zf.read("word/document.xml").decode("utf-8"):
                return False
    return True


def docx_checks() -> dict:
    import zipfile

    from docx import Document
    from docx.oxml.ns import qn

    out = {}
    for name in ("Manuscript_PINO_clean.docx", "Supplementary_Materials_PINO.docx"):
        path = ROOT / "manuscript" / name
        if not path.exists():
            out[name] = {"exists": False}
            continue
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            media = sorted(n for n in zf.namelist() if n.startswith("word/media/"))
        doc = Document(path)
        paras = doc.paragraphs
        txt = "\n".join(p.text for p in paras)
        ref_idx = next(
            (i for i, p in enumerate(paras) if p.text.strip() == "References"),
            len(paras),
        )
        drawing_idx = [i for i, p in enumerate(paras) if p._p.findall(".//" + qn("w:drawing"))]
        latex_cmds = bool(re.search(
            r"\\(frac|mathcal|mathrm|boldsymbol|times|psi|nabla|text|le\b|ge\b|lambda|rho)",
            xml,
        ))
        out[name] = {
            "exists": True,
            "tables": len(doc.tables),
            "omath": len(re.findall(r"<m:oMath[ >]", xml)),
            "omathpara": len(re.findall(r"<m:oMathPara[ >]", xml)),
            "media": len(media),
            "equation_pngs": [m for m in media if re.search(r"eq_\d+", m)],
            "n_drawings": len(drawing_idx),
            "figures_before_references": all(i < ref_idx for i in drawing_idx),
            "raw_dollar": "$" in xml,
            "latex_commands": latex_cmds,
            "equation_or_latex_leak": ("$" in xml) or latex_cmds or ("$$" in txt),
            "refs_61": sum(1 for t in txt.splitlines()
                           if re.match(r"^(?:\[\d+\]|\d+\.)\s", t)) >= 61
            or txt.count("Kasianowicz") >= 1,
        }
    return out


def bibliography_checks() -> dict:
    from docx import Document

    path = ROOT / "manuscript" / "Manuscript_PINO_clean.docx"
    if not path.exists():
        return {"ok": False, "count": 0, "bad": []}
    paras = [p.text.strip() for p in Document(path).paragraphs]
    idx = paras.index("References")
    bib = [p for p in paras[idx + 1:] if re.match(r"^\[\d+\]\s", p)]
    nums = [int(re.match(r"^\[(\d+)\]", p).group(1)) for p in bib]
    bad = [p for p in bib if re.match(r"^\[\d+\]\s+\(\d{4}\)", p)
           or len(p) < 25 or not re.search(r"\(\d{4}\)", p)]
    return {"ok": nums == list(range(1, 63)) and not bad,
            "count": len(bib), "bad": [p[:80] for p in bad][:5]}


def stray_figure_checks() -> dict:
    from PIL import Image

    hits = []
    for path in sorted((ROOT / "scripts" / "figures").glob("*.py")):
        if "_child" in path.read_text():
            hits.append(path.name)
    meta = []
    for path in sorted((ROOT / "manuscript" / "figures").glob("*.png")):
        with Image.open(path) as im:
            for k, v in (im.info or {}).items():
                if "_child" in str(v):
                    meta.append(f"{path.name}:{k}")
    return {"child_refs": hits, "png_meta": meta}



def bracket_integrity_check() -> dict:
    """Every prose square bracket must be a bibliographic citation 1..62."""
    text = (ROOT / "manuscript" / "Manuscript.md").read_text()
    body = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    bad = []
    for m in re.finditer(r"\[([^\]]*)\]", body):
        token = m.group(1)
        if not re.fullmatch(r"\s*\d+(?:\s*(?:--|-|–|,)\s*\d+)*\s*", token):
            bad.append(token)
            continue
        if re.search(r",\S", token):
            bad.append(token)
            continue
        nums = [int(x) for x in re.findall(r"\d+", token)]
        if not nums or min(nums) < 1 or max(nums) > 62:
            bad.append(token)
    typo = bool(re.search(r"28\s*,\s*200", body))
    return {"ok": not bad and not typo, "bad_brackets": bad[:20],
            "bad_count": len(bad), "typo_28_200": typo}


def cpc_and_test_run_checks() -> dict:
    layout = json.loads((ROOT / "results" / "figure_data" /
                         "figure_layout_report.json").read_text())
    specs = {"fig1_workflow_and_architecture": "ABCD",
             "fig2_benchmark_and_field_reconstruction": "ABCDEF",
             "fig3_h3_charge_reparameterization_ablation": "ABCDEF",
             "fig4_h1_h2_interp_vs_ood_regimes": "ABCD",
             "fig5_parametric_physics_sweeps": "ABC",
             "fig6_h4_timing_and_amortization_breakeven": "ABCD"}
    panels_ok = True
    panel_detail = {}
    for name, expected in specs.items():
        info = layout.get("figures", {}).get(name, {}).get("panel_labels", {})
        good = (info.get("expected") == expected and info.get("found") == expected
                and info.get("set_ok") and info.get("all_in_bounds")
                and info.get("ok") and (info.get("min_gap_px") or 0) >= 14.0)
        panels_ok = panels_ok and bool(good)
        panel_detail[name] = info
    man = (ROOT / "manuscript" / "Manuscript.md").read_text()
    labels = [
        "**Program title:**",
        "**CPC Library link to program files:**",
        "**Developer's repository link:**",
        "**Code Ocean capsule:**",
        "**Licensing provisions:**",
        "**Programming language:**",
        "**Supplementary material:**",
        "**Nature of the problem:**",
        "**Solution method:**",
        "**Additional comments including restrictions and unusual features:**",
        "**Running time:**",
    ]
    positions = [man.find(x) for x in labels]
    cpc_ok = all(p >= 0 for p in positions) and positions == sorted(positions)
    cpc_ok = cpc_ok and "Python 3.10+ (tested on 3.14)" in man
    cpc_ok = cpc_ok and "Not applicable" in man and "MIT License (MIT)" in man
    cpc_ok = cpc_ok and "Supplementary Sections S0--S10 and Supplementary Tables S1--S22" in man
    run_files = [PUB / "test_run" / "sample_input.json",
                 PUB / "test_run" / "expected_output.json",
                 PUB / "test_run" / "README.md",
                 PUB / "scripts" / "run_cpc_test_case.py"]
    files_ok = all(p.exists() for p in run_files)
    proc = subprocess.run([sys.executable, "scripts/run_cpc_test_case.py"],
                          cwd=PUB, capture_output=True, text=True, timeout=30) if files_ok else None
    test_ok = bool(proc and proc.returncode == 0 and "CPC TEST RUN PASSED" in proc.stdout)
    return {"panels_ok": bool(panels_ok), "panel_detail": panel_detail,
            "cpc_ok": bool(cpc_ok), "files_ok": files_ok, "test_ok": test_ok,
            "test_output": (proc.stdout.strip().splitlines()[0] if test_ok else
                            (proc.stderr.strip()[-300:] if proc else "missing test_run"))}


def stage16_checks() -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    man = (ROOT / "manuscript" / "Manuscript.md").read_text()
    sup = (ROOT / "manuscript" / "Supplementary_Materials.md").read_text()
    md_caps = [int(m.group(1)) for m in
               re.finditer(r"^\*\*Supplementary Table S(\d+)\.", sup, re.M)]
    doc = Document(ROOT / "manuscript" / "Supplementary_Materials_PINO.docx")
    body_caps: list[int] = []
    tables: list[Table] = []
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            m = re.match(r"^Supplementary Table S(\d+)\.", item.text.strip())
            if m:
                body_caps.append(int(m.group(1)))
        elif isinstance(item, Table):
            tables.append(item)
    header_ok = all(
        t.rows[0]._tr.trPr is not None
        and t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        for t in tables
    )
    no_split_ok = all(
        row._tr.trPr is not None
        and row._tr.trPr.find(qn("w:cantSplit")) is not None
        for t in tables for row in t.rows
    )
    exact_title = (
        "Charge-balanced neural operators, conductance identifiability limits, and warm-started "
        "finite-volume solvers for nanopore ion transport"
    )
    docx_text = "\n".join(p.text for p in doc.paragraphs)
    filename = "scripts/generate_parametric_dataset.py"
    double_abs = [
        re.compile(r"\\ln\s*\|\s*q_w\s*\|\s*\|"),
        re.compile(r"\\ln\s*\\vert\s*q_w\s*\\vert\s*\\vert"),
    ]
    bad_abs = any(p.search(man + "\n" + sup + "\n" + docx_text) for p in double_abs)
    bad_algebra = bool(re.search(
        r"\\sqrt\{s\^2\s*\+\s*\\rho\^2\}\s*=\s*s", man + "\n" + sup))
    return {
        "md_table_order": md_caps,
        "docx_table_order": body_caps,
        "docx_table_count": len(tables),
        "table_order_ok": md_caps == list(range(1, 23)) and body_caps == list(range(1, 23)),
        "header_repeat_ok": header_ok and len(tables) > 0,
        "rows_no_split_ok": no_split_ok and len(tables) > 0,
        "title_ok": man.splitlines()[0].strip() == "# " + exact_title
                    and f"**Manuscript:** {exact_title}" in sup,
        "code_filename_ok": filename in sup and filename in docx_text,
        "algebra_ok": not bad_algebra,
        "double_abs_ok": not bad_abs,
    }


def stage20_checks() -> dict:
    path = ROOT / "results" / "stage20_hybrid_verification.json"
    if not path.exists():
        return {"ok": False, "reason": "missing stage20 verification"}
    d = json.loads(path.read_text())
    corr = d.get("corrector_preserves_reference", {})
    eq = d.get("equal_diffusivity_reversal", {})
    tms = d.get("tms_sign_check", {})
    records = d.get("n_inverse_realizations", 0)
    hybrid = (ROOT / "scripts" / "operators" / "hybrid_controller.py").read_text()
    man = (ROOT / "manuscript" / "Manuscript.md").read_text()
    no_refs = "ref[" not in hybrid
    no_hardcoded = not re.search(r"(?:runtime|wall_time|elapsed)\s*=\s*[0-9]+\.[0-9]+", hybrid)
    no_bad_debye = "k_B T / (2 F^2" not in man and "k_BT/(2F^2" not in man
    no_bad_pct = "54--99%" not in man and "54–99%" not in man
    cases = corr.get("cases", [])
    corr_ok = bool(corr.get("pass")) and bool(cases) and all(
        float(c.get("dpsi", 1e9)) < 1e-9 and float(c.get("dc", 1e9)) < 1e-9
        and float(c.get("current_nonuniformity", 1e9)) < 1e-11 for c in cases)
    eq_ok = bool(eq.get("converged")) and -5.7 < float(eq.get("V_rev_mV", 0)) < -5.6
    tms_ok = abs(float(tms.get("q0_ratio5_mV", 0)) - 0.7254493141463998) < 1e-6
    return {"ok": corr_ok and eq_ok and tms_ok and records == 2700
            and no_refs and no_hardcoded and no_bad_debye and no_bad_pct,
            "corrector": corr_ok, "equal_diffusivity": eq_ok,
            "tms": tms_ok, "records": records,
            "reference_free": no_refs, "no_hardcoded_runtime": no_hardcoded,
            "debye": no_bad_debye, "physics_wording": no_bad_pct}


def stage21_checks() -> dict:
    """Validate the verified Stage-21 evidence: per-case timing arithmetic,
    genuinely populated inverse diagnostics, non-empty bias rows, and a
    Supplementary Table S22 whose caption matches its contents."""
    import numpy as np

    res = ROOT / "results"
    pub = PUB / "results"

    # --- Timing: per-case records plus arithmetic that matches the medians. ---
    timing_path = res / "stage21_timing_benchmark_real.json"
    timing_ok = timing_arith = False
    if timing_path.exists():
        t = json.loads(timing_path.read_text())
        rows = t.get("task_A_states", [])
        s = t.get("summary", {})
        keys = ("warm_sgfv_tol1e-9", "cold_sgfv_tol1e-9", "one_step_gummel")
        timing_ok = (len(rows) == 120
                     and all(all(r.get(k, {}).get("time_s", 0) > 0 for k in keys)
                             for r in rows))
        try:
            ta, tb = s["task_A"], s["task_B"]
            fwd = ta["cold_sgfv_tol1e-9"]["median_time_s"] / ta["warm_sgfv_tol1e-9"]["median_time_s"]
            gum = ta["cold_sgfv_tol1e-9"]["median_time_s"] / ta["one_step_gummel"]["median_time_s"]
            rev = tb["cold_brent_vrev"]["median_time_s"] / tb["warm_secant_vrev"]["median_time_s"]
            timing_arith = (
                abs(fwd - s["task_A_speedups"]["warm_1e-9_vs_cold_1e-9"]) <= 1e-6 * max(1.0, fwd)
                and abs(gum - s["task_A_speedups"]["one_step_vs_cold_1e-9"]) <= 1e-6 * max(1.0, gum)
                and abs(rev - s["task_B_speedups"]["warm_vs_cold"]) <= 1e-6 * max(1.0, rev)
            )
        except (KeyError, ZeroDivisionError):
            timing_arith = False

    # --- Inverse: per-realization SEs/intervals/flags, not placeholders. ---
    inv_path = res / "stage21_inverse_records_real.npz"
    sum_path = res / "stage21_inverse_summary_real.json"
    inv_ok = diag_ok = cov_ok = coverage_matches = False
    n_records = 0
    if inv_path.exists():
        z = np.load(inv_path, allow_pickle=True)
        n_records = len(z["B_se_lnq"]) if "B_se_lnq" in z.files else 0
        se_fields = [f"{e}_se_{q}" for e in "ABC" for q in ("lnR", "lnsigma", "lnq")]
        have = [f for f in se_fields if f in z.files]
        diag_ok = (n_records >= 2600 and len(have) == len(se_fields)
                   and all(np.mean(np.isfinite(z[f]) & (z[f] > 0)) > 0.99 for f in have)
                   and "B_delta_m_inf" in z.files
                   and np.mean(np.isfinite(z["B_delta_m_inf"]) & (z["B_delta_m_inf"] > 0)) > 0.99)
        cov_fields = [f"{e}_{q}_cover" for e in "ABC" for q in ("q", "R", "sigma")]
        cov_fields = [f for f in cov_fields if f in z.files]
        cov_ok = len(cov_fields) == 9 and all(
            np.all(np.isfinite(z[f].astype(float))) and np.all((z[f].astype(float) >= 0) & (z[f] <= 1))
            for f in cov_fields)
        inv_ok = diag_ok and cov_ok
        if sum_path.exists() and inv_ok:
            sj = json.loads(sum_path.read_text())
            proto = np.array([p.decode() if isinstance(p, bytes) else str(p) for p in z["protocol"]])
            noise = z["noise"].astype(float)
            m = (proto == "dual") & (np.abs(noise - 0.01) < 1e-12)
            if m.any():
                rec_cov = float(np.mean(z["B_q_cover"][m].astype(float)))
                summary_cov = float(sj["by_condition"]["dual_0.01"]["B_q_cover"])
                coverage_matches = abs(rec_cov - summary_cov) <= 1e-12
                inv_ok = inv_ok and coverage_matches

    # --- Bias: rows must exist for D1/D2 and D3 must summarise real records. ---
    bias_path = res / "stage21_bias_diagnostic_real.json"
    bias_ok = False
    if bias_path.exists():
        b = json.loads(bias_path.read_text())
        rows = b.get("rows", [])
        labels = {r.get("diagnostic") for r in rows}
        bias_ok = (len(rows) >= 30
                   and {"same_grid_noise_free", "cross_grid_noise_free"} <= labels
                   and b.get("summary", {}).get("diagnostic_3_warm_vs_cold_noisy", {}).get("n", 0) > 0)

    # --- Published copies and the claimed file set. ---
    pub_files = ["stage21_timing_benchmark_real.json",
                 "stage21_inverse_records_real.npz",
                 "stage21_inverse_summary_real.json",
                 "stage21_bias_diagnostic_real.json"]
    pub_ok = all((pub / f).exists() and (pub / f).stat().st_size > 10000 for f in pub_files)

    # --- Table S22 caption must match displayed content. ---
    sup = (ROOT / "manuscript" / "Supplementary_Materials.md").read_text()
    s22 = sup[sup.index("**Supplementary Table S22."):] if "**Supplementary Table S22." in sup else ""
    s22_ok = ("coverage" in s22.lower() and "sigma" in s22.lower()
              and all(tag in s22 for tag in ("D1", "D2", "D3")))

    man = (ROOT / "manuscript" / "Manuscript.md").read_text()
    cap = ""
    for para in man.split("\n\n"):
        if para.startswith("**Figure 1.**"):
            cap = para
            break
    cap_ok = ("zero wall ion flux, the wall surface-charge condition, and interior PDE residuals are enforced through the physics loss" in cap
              and "ansatz imposes reservoir potential and concentration values, zero wall ion flux" not in cap)

    ok = (timing_ok and timing_arith and inv_ok and bias_ok and pub_ok and s22_ok and cap_ok)
    return {"ok": ok, "public_files": pub_ok,
            "timing": timing_ok, "timing_arithmetic": timing_arith,
            "inverse": inv_ok, "inverse_diagnostics": diag_ok,
            "inverse_coverage": cov_ok, "coverage_matches_summary": coverage_matches,
            "n_records": n_records, "bias": bias_ok,
            "figure1_caption": cap_ok, "s22": s22_ok}


def main() -> None:
    tokens, apis_names = secret_tokens()
    public = scan_tree(PUB, tokens)
    private = scan_tree(PRIV, tokens)
    unignored = scan_git_unignored(tokens)
    ms = manuscript_checks()
    claims = claim_checks()
    display = display_item_checks()
    docx = docx_checks()
    seq = sequential_checks()
    s11 = stage11_checks()
    bib = bibliography_checks()
    stray = stray_figure_checks()
    brackets = bracket_integrity_check()
    cpc = cpc_and_test_run_checks()
    s16 = stage16_checks()
    s20 = stage20_checks()
    s21 = stage21_checks()
    summary = json.loads((ROOT / "reproducibility" / "PACKAGE_SUMMARY.json").read_text())

    def ok(cond: bool) -> str:
        return "PASS" if cond else "FAIL"

    checks = [
        ("No credential values in public package", ok(not public["secret_value_hits"])),
        ("No credential values in private package", ok(not private["secret_value_hits"])),
        ("No credential values in tracked/unignored project files", ok(not unignored["secret_value_hits"])),
        ("No secret file/dir names in public package", ok(not public["secret_name_hits"])),
        ("No secret file/dir names in private package", ok(not private["secret_name_hits"])),
        ("No generic key patterns in public package", ok(not public["pattern_hits"])),
        ("No generic key patterns in private package", ok(not private["pattern_hits"])),
        ("Abstract <= 250 words", ok(0 < ms["abstract_words"] <= 250)),
        ("3-5 highlights, each <= 85 chars", ok(3 <= ms["highlights"] <= 5 and ms["highlights_max_chars"] <= 85)),
        ("All cited reference numbers resolve", ok(not ms["cited_missing"])),
        ("Every reference is cited", ok(not ms["uncited"])),
        ("Every reference has a DOI", ok(ms["refs_with_doi"] == ms["references"])),
        ("No hidden placeholders remain", ok(not ms["placeholders"])),
        ("Dimensionless system in supplement", ok(ms["supp_has_dimensionless_section"])),
        ("Exponential reservoir ansatz documented", ok(ms["supp_ansatz_exponential"])),
        ("Claim-evidence map complete", ok(claims["all_evidenced"])),
        ("Main DOCX present, no LaTeX leak", ok(docx["Manuscript_PINO_clean.docx"].get("exists") and not docx["Manuscript_PINO_clean.docx"]["equation_or_latex_leak"])),
        ("Supplementary DOCX present, no LaTeX leak", ok(docx["Supplementary_Materials_PINO.docx"].get("exists") and not docx["Supplementary_Materials_PINO.docx"]["equation_or_latex_leak"])),
        ("Main DOCX: native OMML equations, no equation images, figures inline before References",
         ok(docx["Manuscript_PINO_clean.docx"].get("exists")
            and docx["Manuscript_PINO_clean.docx"]["omath"] > 50
            and docx["Manuscript_PINO_clean.docx"]["tables"] == 0
            and not docx["Manuscript_PINO_clean.docx"]["equation_pngs"]
            and docx["Manuscript_PINO_clean.docx"]["media"] == 6
            and docx["Manuscript_PINO_clean.docx"]["figures_before_references"])),
        ("Supplementary DOCX: native OMML equations, no images",
         ok(docx["Supplementary_Materials_PINO.docx"].get("exists")
            and docx["Supplementary_Materials_PINO.docx"]["omath"] > 50
            and docx["Supplementary_Materials_PINO.docx"]["media"] == 0)),
        ("Public package has no private logs/download/literature", ok(not any((PUB / p).exists() for p in ("logs", "download", "literature", "apis", ".git")))),
        ("Private package contains manuscript and reports", ok((PRIV / "manuscript/Manuscript.md").exists() and (PRIV / "reports/STATUS.md").exists())),
        ("Six figures cited sequentially; no main-text data tables; 22 supplement tables cited",
         ok(display["figure_order"] == [1, 2, 3, 4, 5, 6] and display["body_tables"] == 0
            and display["supp_tables"] == 22 and not display["uncited_supp_tables"])),
        ("Claim-evidence map spans C01--C44 and L01--L62, all resolved",
         ok(claims["quant_complete"] and claims["lit_complete"] and claims["all_resolved"])),
        ("References are numbered strictly by first appearance 1..62",
         ok(seq["ref_order"] == list(range(1, 63)))),
        ("Supplementary tables and equations are strictly sequential",
         ok(seq["tab_order"] == list(range(1, 23))
            and seq["eq_main"]["strict_sequential"] and seq["eq_supp"]["strict_sequential"])),
        ("Figures are placed inline within 1-2 paragraphs of the first mention (Markdown + DOCX)",
         ok(all(v["image_after_mention"] and v["caption_after_image"] for v in seq["figs"].values())
            and seq["figs_docx"].get("all_inline", False))),
        ("Stage-11 reviewer upgrades present (tables, JSON, statistics, script)",
         ok(s11["tables"] and all(s11["json_keys"].values()) and s11["statistics_keys"] and s11["script"])),
        ("DOCX bibliography is numbered [1]--[62] with authors and years",
         ok(bib["ok"])),
        ("Supplementary tables first-seen 1..22; no stray legend children in figures",
         ok(seq["tab_order"] == list(range(1, 23))
            and not stray["child_refs"] and not stray["png_meta"])),
        ("Official Zenodo DOI present in manuscript, main DOCX, CPC letter, and public README",
         ok(zendo_doi_check())),
        ("Zero non-citation square brackets in prose; no 28,200 typo",
         ok(brackets["ok"])),
        ("Figure 1--6 panel labels present/spaced; official CPC template and test_run pass",
         ok(cpc["panels_ok"] and cpc["cpc_ok"] and cpc["files_ok"] and cpc["test_ok"])),
        ("Supplementary tables physically ordered S1--S22; Word headers repeat and rows cannot split",
         ok(s16["table_order_ok"] and s16["header_repeat_ok"] and s16["rows_no_split_ok"])),
        ("Final title, algebra, code filename, and single-absolute Fisher notation",
         ok(s16["title_ok"] and s16["algebra_ok"] and s16["code_filename_ok"]
            and s16["double_abs_ok"])),
        ("Stage-20 corrector, equal-diffusivity reversal, TMS sign, blind 2,700-realization inverse, and honest physics wording",
         ok(s20["ok"])),
        ("Stage-21 raw records published, equal-workload timing, 3-way bias diagnostic, Figure-1 caption, and separate R/sigma/q coverage in S22",
         ok(s21["ok"])),
    ]

    lines = [
        "# Final audit report",
        "",
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')}",
        "",
        "Audited packages:",
        f"- Public: {summary['public_repo']['file_count']} files, "
        f"{summary['public_repo']['total_bytes'] / 1e6:.1f} MB",
        f"- Private: {summary['private_repo']['file_count']} files, "
        f"{summary['private_repo']['total_bytes'] / 1e6:.1f} MB",
        "",
        "## Checklist",
        "",
        "| Check | Result |",
        "|---|---|",
    ]
    lines += [f"| {name} | {res} |" for name, res in checks]
    lines += [
        "",
        "## Details",
        "",
        f"- Credential files inspected: {len(apis_names)} (values never printed).",
        f"- Public package secret-value hits: {public['secret_value_hits']}",
        f"- Private package secret-value hits: {private['secret_value_hits']}",
        f"- Tracked/unignored files scanned: {unignored['files']}; hits: {unignored['secret_value_hits']}",
        f"- Public secret-name hits: {public['secret_name_hits']}",
        f"- Private secret-name hits: {private['secret_name_hits']}",
        f"- Abstract: {ms['abstract_words']} words; highlights: {ms['highlights']} "
        f"(max {ms['highlights_max_chars']} chars).",
        f"- References: {ms['references']}, cited: {ms['cited']}, "
        f"uncited: {ms['uncited']}, missing: {ms['cited_missing']}, with DOI: {ms['refs_with_doi']}.",
        f"- Claims: {claims['claims']} ({claims['quant']} quantitative + {claims['lit']} literature), "
        f"statuses: {claims['statuses']}.",
        f"- Display items: figure first-citation order {display['figure_order']}; "
        f"main-text data tables {display['body_tables']}; supplement tables {display['supp_tables']}; "
        f"uncited supplement tables {display['uncited_supp_tables']}.",
        f"- Sequential numbering: references {seq['ref_order'][:5]}... (n={len(seq['ref_order'])}); "
        f"supplementary tables {seq['tab_order']}; "
        f"main equations {seq['eq_main']['first_appearance']}; "
        f"supplement equations {seq['eq_supp']['first_appearance']}.",
        f"- Stage-11 upgrades: {json.dumps(s11)}.",
        f"- Prose bracket integrity: {json.dumps(brackets)}.",
        f"- CPC template/test run: {json.dumps(cpc)}.",
        f"- DOCX: {json.dumps(docx)}.",
        "",
        "## Scope and limitations",
        "",
        "- The secret scan checks exact credential values, common key prefixes, and",
        "  secret file names. It does not prove the absence of unknown secrets.",
        "- Citation resolution checks numbering and DOI presence, not whether each",
        "  source supports the associated claim; claim support is tracked separately",
        "  in `manuscript/claim_evidence.csv`.",
        "- The DOCX was built from the canonical markdown; rendered-page inspection",
        "  is limited to structural checks (tables, embedded figures, equation images,",
        "  and absence of LaTeX leakage).",
        "- The archived DOI is the official Zenodo record 10.5281/zenodo.23487404,",
        "  checked in the manuscript, main DOCX, CPC cover letter, and public README.",
    ]
    REPORT.write_text("\n".join(lines) + "\n")
    print("\n".join(f"{res:4} {name}" for name, res in checks))
    print(f"report: {REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
